# coding=utf-8
from win32com import client as wc
from docx import Document
import xlwt
import re
import os


def docx_to_dict(path):
    print("提取文档中景区信息...{}".format(path))
    document = Document(path)
    l = [paragraph.text for paragraph in document.paragraphs]

    scenic = {
        'name': '',
        'level': '',
        'pinyin': '',
        'spelling': '',
        'date': '',
        'tel': '',
        'country': '',
        'province': '',
        'city': '',
        'region': '',
        'address': '',
        'intro': '',
        'spot': []
    }

    # ---------------------------------pattern-------------------------------------------
    name_pat = re.compile(r'(.*?)（.*?字?数?.*?）')
    level_pat = re.compile(r'(.*?)级?景[点区]|([无未][等评]级)|(\dA)级')
    pinyin_pat = re.compile(r'^\s*([a-zA-Z]+)\s*')
    date_pat = re.compile(r'(开放时间|夏季|旺季):?\s*(.*)')
    tel_pat = re.compile(r'(电话|方式):?\s*(.*)')
    addr_pat = re.compile(r'地址:\s*(.*)')
    none_pat = re.compile(r'^$')
    # spot_name_pat = re.compile(r'【\d\d\s*(.*?)】')
    spot_name_pat = re.compile(r'【\d\d\s*(.*?)】.*?\n', re.S)
    end_pat = re.compile(r'【?以下.*?[录配]音\s*】?|【交通建议】')

    # ------------------------------Pre-treatment----------------------------------------
    content = ''
    for phr in l:
        phr = phr.replace('：', ':').replace('(', '（').replace(')', '）')
        if phr:
            content += '\n' + phr
    spot_start = 0
    spot_end = 0
    try:
        spot_start = spot_name_pat.search(content).regs[0][0]
        spot_end = end_pat.search(content).regs[0][0] - 1

    except Exception:
        print("文件预处理出错\t{0}".format(path))
    main_info = content[:spot_start - 1]
    spot_info = content[spot_start:spot_end]
    # ------------------------------------------- main-info-----------------------------------------------
    for phr in main_info.split('\n'):
        if not phr:
            continue
        if name_pat.findall(phr):
            scenic['name'] = name_pat.findall(phr)[0]
            name_pat = none_pat
        elif level_pat.findall(phr):
            scenic['level'] = level_pat.findall(phr)[0]
            scenic['level'] = ''.join(scenic['level'])
            level_pat = none_pat
        elif pinyin_pat.findall(phr):  # 根据长短进行简拼和全拼判断
            if not scenic['pinyin']:
                scenic['pinyin'] = pinyin_pat.findall(phr)[0]
            else:
                scenic['spelling'] = pinyin_pat.findall(phr)[0]
                if len(scenic['pinyin']) < len(scenic['spelling']):
                    scenic['pinyin'], scenic['spelling'] = scenic['spelling'], scenic['pinyin']

        elif date_pat.findall(phr):
            scenic['date'] = date_pat.findall(phr)[0][1]
            date_pat = none_pat
        elif tel_pat.findall(phr):
            scenic['tel'] = tel_pat.findall(phr)[0][1]
            tel_pat = none_pat
        elif addr_pat.findall(phr):
            scenic['address'] = addr_pat.findall(phr)[0]
            addr_pat = none_pat
        else:
            print("未找到合适匹配\t{1}\t{0})".format(path, phr.split()))

    # ------------------------------------------- spot-info-----------------------------------------------
    spot_names = spot_name_pat.split(spot_info)
    spot_names.pop(0)
    spot_names.pop(0)

    scenic['intro'] = spot_names.pop(0)
    while spot_names:
        scenic['spot'].append({'name': spot_names.pop(0), 'desc': spot_names.pop(0)})

    return scenic


def doc_to_docx(src, des):
    print("进行doc->docx文件格式转换...{}".format(src))
    word = wc.Dispatch('Word.Application')
    doc = word.Documents.Open(src)  # 目标路径下的文件
    doc.SaveAs(des, 16)  # 转化后路径下的文件
    doc.Close()
    word.Quit()
    return des


def dict_to_excel(info):
    workbook = xlwt.Workbook(encoding='ascii')
    worksheet = workbook.add_sheet('scenic')
    worksheet.write(0, 0, label='Row 0, Column 0 Value')
    workbook.save('Excel_Workbook.xls')


def main():
    docx_files = []
    for root, dirs, files in os.walk(r'E:\Github\python\data\5.28文稿'):
        for name in files:
            suffix = os.path.splitext(name)[1]
            if suffix == '.docx' and '$' not in name:
                docx_files.append(os.path.join(root, name))
            if suffix == '.doc' and '$' not in name:
                doc_path = os.path.join(root, name)
                des = 'E:/Github/python/data/convert_docx/' + name.replace('.doc', '.docx')
                try:
                    docx_path = doc_to_docx(doc_path.replace('\\', '/'), des)
                    docx_files.append(docx_path)
                except Exception:
                    print("格式转换(doc-docx)出错\t{0}".format(doc_path))

    workbook = xlwt.Workbook(encoding='ascii')
    worksheet = workbook.add_sheet('scenic')
    worksheet.write(0, 0, label='景区名称')
    worksheet.write(0, 1, label='景区级别')
    worksheet.write(0, 2, label='全拼')
    worksheet.write(0, 3, label='简拼')
    worksheet.write(0, 4, label='开放时间')
    worksheet.write(0, 5, label='联系电话')
    worksheet.write(0, 6, label='国家')
    worksheet.write(0, 7, label='省')
    worksheet.write(0, 8, label='市')
    worksheet.write(0, 9, label='区')
    worksheet.write(0, 10, label='地址')
    worksheet.write(0, 11, label='景区文本')

    for row, file in enumerate(docx_files):
        scenic_info = docx_to_dict(file)
        column = 0
        for key, value in scenic_info.items():
            if key != 'spot':
                worksheet.write(row + 1, column, label=value)
                column += 1

            else:
                for column2, spot in enumerate(value):
                    try:

                        worksheet.write(0, column + column2 * 2, label='景点' + str(column2 + 1) + '名称')
                        worksheet.write(0, column + column2 * 2 + 1, label='景点' + str(column2 + 1) + '简介')
                    except Exception:
                        pass
                    worksheet.write(row + 1, column + column2 * 2, label=spot['name'])
                    worksheet.write(row + 1, column + column2 * 2 + 1, label=spot['desc'])

        workbook.save('Excel_Workbook.xls')


if __name__ == "__main__":
    main()
